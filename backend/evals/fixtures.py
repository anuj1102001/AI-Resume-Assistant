"""Render the synthetic resumes to the file formats people actually upload (PDF, two-column PDF, DOCX)."""

import html
import io
import json
from pathlib import Path

import pymupdf

from app.resume.sections import heading_key

DATA = Path(__file__).parent / "data"
RESUMES = DATA / "resumes"
JOBS = DATA / "jobs"

# How each resume is rendered. Rohan's is two columns on purpose: it's the classic layout that breaks parsers.
FORMATS = {"rohan": "pdf-2col", "kavya": "docx"}


def _html(lines: list[str]) -> str:
    out = []
    for i, line in enumerate(lines):
        esc = html.escape(line)
        if i == 0:
            out.append(f"<p style='font-size:16pt;font-weight:bold;margin:0 0 4pt 0'>{esc}</p>")
        elif heading_key(line):
            out.append(f"<p style='font-size:11pt;font-weight:bold;margin:10pt 0 3pt 0'>{esc}</p>")
        else:
            out.append(f"<p style='font-size:9.5pt;margin:0 0 2pt 0'>{esc}</p>")
    return f"<div style='font-family:sans-serif'>{''.join(out)}</div>"


def _pdf(columns: list[list[str]]) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    if len(columns) == 1:
        page.insert_htmlbox(pymupdf.Rect(48, 48, 547, 800), _html(columns[0]))
    else:
        page.insert_htmlbox(pymupdf.Rect(40, 48, 220, 800), _html(columns[0]))
        page.insert_htmlbox(pymupdf.Rect(236, 48, 555, 800), _html(columns[1]))
    data = doc.tobytes()
    doc.close()
    return data


def _docx(lines: list[str]) -> bytes:
    import docx

    d = docx.Document()
    for i, line in enumerate(lines):
        if i == 0:
            d.add_heading(line, 0)
        elif heading_key(line):
            d.add_heading(line, 1)
        else:
            d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render(name: str) -> tuple[bytes, str]:
    lines = [ln for ln in (RESUMES / f"{name}.txt").read_text(encoding="utf-8").split("\n") if ln.strip()]
    fmt = FORMATS.get(name, "pdf")
    if fmt == "docx":
        return _docx(lines), f"{name}.docx"
    if fmt == "pdf-2col":
        # Left column: name, contact and the short sections. Right column: experience.
        start = next(i for i, ln in enumerate(lines) if heading_key(ln) == "experience")
        end = next(i for i, ln in enumerate(lines) if i > start and heading_key(ln))
        return _pdf([lines[:start] + lines[end:], lines[start:end]]), f"{name}.pdf"
    return _pdf([lines]), f"{name}.pdf"


def names() -> list[str]:
    return sorted(p.stem for p in RESUMES.glob("*.txt"))


def gold() -> dict:
    return json.loads((DATA / "resumes.json").read_text(encoding="utf-8"))


def job_text(name: str) -> str:
    return (JOBS / f"{name}.txt").read_text(encoding="utf-8")
