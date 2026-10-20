"""Bullet rewriting with guardrails, tailored versions, and export.

The model is told not to add facts, but we check anyway. Every variant is compared with the original:
  new_numbers  numbers in the rewrite that aren't in the original (placeholders like [X]% are fine and expected)
  new_skills   skills/tools the rewrite names that appear nowhere in the resume
Flagged variants are shown with a warning; the user decides what to keep.
"""

import io
import re

from app.db import Database, loads, new_id, now
from app.errors import AppError
from app.grounding import numbers_in
from app.llm import prompts
from app.llm.client import LLM
from app.llm.schemas import RewriteOut
from app.resume.skills import find_skills
from app.services import grounded, jobs, resumes

_PLACEHOLDER = re.compile(r"\[[^\]]*\]")


def validate_variant(variant: str, original: str, resume_text: str) -> dict:
    stripped = _PLACEHOLDER.sub("", variant)
    new_numbers = sorted(numbers_in(stripped) - numbers_in(original))
    resume_skills = set(find_skills(resume_text))
    new_skills = sorted(set(find_skills(stripped)) - resume_skills)
    return {
        "new_numbers": new_numbers,
        "new_skills": new_skills,
        "placeholders": _PLACEHOLDER.findall(variant),
        "words": len(variant.split()),
    }


def _context(parsed: dict, bullet_id: str) -> tuple[str, str]:
    bullet = next((b for b in parsed["bullets"] if b["id"] == bullet_id), None)
    if not bullet:
        raise AppError("That bullet isn't in this resume.", 404, "NOT_FOUND")
    siblings = [
        b["text"] for b in parsed["bullets"] if b["entry"] == bullet["entry"] and b["section"] == bullet["section"] and b["id"] != bullet_id
    ]
    header = (
        parsed["entries"][bullet["entry"]]["header"] if bullet["entry"] is not None and bullet["entry"] < len(parsed["entries"]) else ""
    )
    return bullet["text"], "\n".join([header, *siblings[:5]]).strip()


def rewrite(db: Database, llm: LLM, ws_id: str, rid: str, bullet_id: str, jid: str | None, refresh: bool = False) -> dict:
    resume = resumes.get_resume(db, ws_id, rid)
    job = jobs.get_job(db, ws_id, jid) if jid else None
    parsed = loads(resume["parsed"], {})
    original, context = _context(parsed, bullet_id)
    key = grounded.cache_key("rewrite", [rid, bullet_id, jid or ""], llm.model)
    if not refresh and (hit := grounded.cached(db, ws_id, key)):
        return hit
    out, usage = llm.generate(
        system=prompts.SYSTEM,
        prompt=prompts.rewrite_prompt(original, context, job["text"] if job else None),
        schema=RewriteOut,
        temperature=0.6,
    )
    variants = [{**v.model_dump(), "checks": validate_variant(v.text, original, resume["text"])} for v in out.variants]
    result = {"bullet_id": bullet_id, "original": original, "variants": variants, "questions": out.questions}
    return grounded.save(db, ws_id, "rewrite", key, rid, jid, {"bullet_id": bullet_id}, result, usage)


def save_tailored(db: Database, ws_id: str, rid: str, jid: str | None, bullet_id: str, revised: str) -> dict:
    resume = resumes.get_resume(db, ws_id, rid)
    if jid:
        jobs.get_job(db, ws_id, jid)
    original, _ = _context(loads(resume["parsed"], {}), bullet_id)
    revised = revised.strip()
    if not revised:
        raise AppError("The revised bullet is empty.", 422, "INVALID_INPUT")
    with db.tx() as c:
        c.execute("DELETE FROM tailored WHERE resume_id = ? AND COALESCE(job_id, '') = ? AND original = ?", (rid, jid or "", original))
        if revised != original:
            c.execute(
                "INSERT INTO tailored (id, workspace_id, resume_id, job_id, original, revised, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (new_id(), ws_id, rid, jid, original, revised, now()),
            )
    return {"bullet_id": bullet_id, "original": original, "revised": revised, "checks": validate_variant(revised, original, resume["text"])}


def tailored_version(db: Database, ws_id: str, rid: str, jid: str | None) -> dict:
    resume = resumes.get_resume(db, ws_id, rid)
    rows = db.all(
        "SELECT original, revised, created_at FROM tailored WHERE resume_id = ? AND COALESCE(job_id, '') = ? ORDER BY created_at",
        (rid, jid or ""),
    )
    changes = {r["original"]: r["revised"] for r in rows}
    parsed = loads(resume["parsed"], {})
    bullets = [
        {"id": b["id"], "section": b["section"], "original": b["text"], "revised": changes.get(b["text"])} for b in parsed["bullets"]
    ]
    return {"resume_id": rid, "job_id": jid, "changed": sum(1 for b in bullets if b["revised"]), "bullets": bullets}


def export_docx(db: Database, ws_id: str, rid: str, jid: str | None) -> tuple[bytes, str]:
    """A clean, single-column, ATS-friendly DOCX of the resume with accepted rewrites applied."""
    import docx
    from docx.shared import Pt

    resume = resumes.get_resume(db, ws_id, rid)
    parsed = loads(resume["parsed"], {})
    changes = {b["original"]: b["revised"] for b in tailored_version(db, ws_id, rid, jid)["bullets"] if b["revised"]}
    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10.5)
    for sec in parsed["sections"]:
        if sec["key"] == "header":
            lines = sec["text"].split("\n")
            if lines:
                d.add_heading(lines[0], level=0)
                for line in lines[1:]:
                    d.add_paragraph(line)
            continue
        d.add_heading(sec["title"] or resumes.SECTION_LABEL.get(sec["key"], sec["key"].title()), level=1)
        section_bullets = [b for b in parsed["bullets"] if b["section"] == sec["key"]]
        pending = list(section_bullets)
        for line in sec["text"].split("\n"):
            text = re.sub(r"^\s*[•●▪■◦‣∙·–—*➢➤►-]\s*", "", line).strip()
            if not text or len(text) == 1:
                continue
            if pending and pending[0]["text"].startswith(text[:30]):
                b = pending.pop(0)
                d.add_paragraph(changes.get(b["text"], b["text"]), style="List Bullet")
            elif any(text in b["text"] for b in section_bullets):
                continue  # the wrapped tail of a bullet already written
            else:
                d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    name = re.sub(r"[^A-Za-z0-9 _-]", "", resume["label"]).strip().replace(" ", "_") or "resume"
    return buf.getvalue(), f"{name}{'_tailored' if changes else ''}.docx"
